"""Global sea level through time, from four reconstructions shown side by side.

Each series is metres above (+) or below (-) today's sea level, by age in thousands of years
(ka). They differ in method, span and how much they smooth, and they disagree, which the site
shows rather than averages away:

- van der Meer et al. (2022), Gondwana Research 111: 0-540 Ma every million years, with a
  range: plate-tectonic sea level from strontium isotopes, plus ice-sheet volume (CC BY 4.0,
  Supplementary Table 1);
- Miller et al. (2020), Science Advances 6: 1-65 Ma every 20,000 years, from deep-sea
  oxygen isotopes and Mg/Ca, smoothed (PANGAEA, CC BY 4.0);
- Bintanja & van de Wal (2008), Nature 454: 0-3 Ma, from an ice-sheet model driven by the
  LR04 oxygen-isotope stack (NOAA; thinned from 100-year to 1,000-year steps). Its "RSL"
  column is the fall below today (+123 m at the last glacial maximum), so its sign is turned;
- Spratt & Lisiecki (2016), Climate of the Past 12: 0-798 ka, a stack of five records with a
  95% range (NOAA). Its 0 ka value is about +9 m: the stack's zero is not today's sea level.

Haq's widely drawn curves are left out: their underlying data are not published.
"""

import io
import json
import urllib.request
from dataclasses import dataclass
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from histeclipse_pipeline.sky import Json, verified

if TYPE_CHECKING:
    from pathlib import Path


@dataclass(frozen=True)
class Source:
    """Where a series comes from, and the file it is read from."""

    url: str
    sha256: str


VAN_DER_MEER = Source(
    "https://ars.els-cdn.com/content/image/1-s2.0-S1342937X22002192-mmc1.xlsx",
    "800b34928378ae902426fc16ef002debfc064465f8ac663cbeeb7bd0c77c074f",
)
MILLER = Source(
    "https://doi.pangaea.de/10.1594/PANGAEA.923139?format=textfile",
    "fdd0f1d5ff5863b8eb72d47ddd98cfe6464c0ffa83d97e0f8a0597233c773796",
)
BINTANJA = Source(
    "https://www.ncei.noaa.gov/pub/data/paleo/contributions_by_author/bintanja2008/bintanja2008-noaa.txt",
    "2f67a3e57034e0b59a50a661adbdbbf3b77c45e44a6efb0917986982319353f5",
)
SPRATT = Source(
    "https://www.ncei.noaa.gov/pub/data/paleo/contributions_by_author/spratt2016/spratt2016-noaa.txt",
    "02f56f02f709b70ac70fe5e403326272e19b42bf34ee9b1bba3b8322e1375dc1",
)

#: van der Meer et al.'s table: age in Ma, and the tectono-glacio-eustatic curve's columns.
_VDM_AGE = 0
_VDM_LEVEL, _VDM_LOW, _VDM_HIGH = 34, 35, 36
_VDM_FIRST_ROW = 5


@dataclass(frozen=True)
class Series:
    """One reconstruction: ages (ka) and sea level (m), with a range where the source gives one."""

    id: str
    label: str
    method: str
    source: str
    url: str
    licence: str
    note: str | None
    ages: list[float]
    level: list[float]
    low: list[float] | None = None
    high: list[float] | None = None


def noaa_table(text: str) -> list[dict[str, str]]:
    """Read a NOAA paleo template file: `#` lines of metadata, then a tab-separated table."""
    lines = [line for line in text.splitlines() if line.strip() and not line.startswith("#")]
    header, *rows = (line.split("\t") for line in lines)
    return [dict(zip(header, row, strict=True)) for row in rows]


def pangaea_table(text: str) -> list[list[str]]:
    """Read a PANGAEA text export: a `/* ... */` description, a header line, then rows."""
    _, _, body = text.partition("*/")
    lines = [line for line in body.splitlines() if line.strip()]
    return [line.split("\t") for line in lines[1:]]


def spratt(text: str) -> Series:
    """Spratt & Lisiecki (2016): the five-record stack over 0-798 ka, with its 95% range."""
    rows = [row for row in noaa_table(text) if row["SeaLev_longPC1"].strip() not in {"", "NaN"}]
    return Series(
        id="spratt-2016",
        label="Spratt & Lisiecki 2016",
        method="stack of five sea-level records from deep-sea oxygen isotopes",
        source="Spratt & Lisiecki 2016, Climate of the Past 12; NOAA",
        url="https://doi.org/10.5194/cp-12-1079-2016",
        licence="NOAA World Data Service; cite the publication",
        note="The stack reads about +9 m at 0 ka: its zero is not today's sea level.",
        ages=[float(row["age_calkaBP"]) for row in rows],
        level=[round(float(row["SeaLev_longPC1"]), 1) for row in rows],
        low=[round(float(row["SeaLev_longPC1_err_lo"]), 1) for row in rows],
        high=[round(float(row["SeaLev_longPC1_err_up"]), 1) for row in rows],
    )


def bintanja(text: str) -> Series:
    """Bintanja & van de Wal (2008): 0-3 Ma, every 1,000 years of the 100-year model output."""
    rows = [row for row in noaa_table(text) if float(row["age_calkaBP"]).is_integer()]
    return Series(
        id="bintanja-2008",
        label="Bintanja & van de Wal 2008",
        method="ice-sheet model driven by the LR04 oxygen-isotope stack",
        source="Bintanja & van de Wal 2008, Nature 454; NOAA",
        url="https://doi.org/10.1038/nature07158",
        licence="NOAA World Data Service; cite the publication",
        note="The source gives the fall below today; it is shown here as height above.",
        ages=[float(row["age_calkaBP"]) for row in rows],
        # RSL counts a fall as positive (+123 m at 20 ka, matching the ice volume).
        level=[round(-float(row["RSL"]), 1) + 0.0 for row in rows],
    )


def miller(text: str) -> Series:
    """Miller et al. (2020): the smoothed Cenozoic curve, 1-65 Ma every 20,000 years."""
    rows = pangaea_table(text)
    return Series(
        id="miller-2020",
        label="Miller et al. 2020",
        method="deep-sea oxygen isotopes and Mg/Ca, checked against continental margins",
        source="Miller et al. 2020, Science Advances 6; PANGAEA",
        url="https://doi.org/10.1594/PANGAEA.923139",
        licence="CC BY 4.0",
        note="Smoothed. The authors flag Mg/Ca records older than 48 Ma as less certain.",
        ages=[float(age) for age, _ in rows],
        level=[round(float(level), 1) for _, level in rows],
    )


def van_der_meer(data: bytes) -> Series:
    """Van der Meer et al. (2022): the tectono-glacio-eustatic curve, 0-540 Ma, with its range."""
    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    ages: list[float] = []
    level: list[float] = []
    low: list[float] = []
    high: list[float] = []
    for row in sheet.iter_rows(min_row=_VDM_FIRST_ROW, values_only=True):
        cells = [cell for cell in row if isinstance(cell, int | float)]
        if len(cells) != len(row):
            continue
        ages.append(float(cells[_VDM_AGE]) * 1000)
        level.append(round(float(cells[_VDM_LEVEL]), 1))
        low.append(round(float(cells[_VDM_LOW]), 1))
        high.append(round(float(cells[_VDM_HIGH]), 1))
    return Series(
        id="van-der-meer-2022",
        label="van der Meer et al. 2022",
        method="plate-tectonic sea level from strontium isotopes, plus ice-sheet volume",
        source="van der Meer et al. 2022, Gondwana Research 111, Supplementary Table 1",
        url="https://doi.org/10.1016/j.gr.2022.07.014",
        licence="CC BY 4.0",
        note=None,
        ages=ages,
        level=level,
        low=low,
        high=high,
    )


def series_json(series: Series) -> dict[str, Json]:
    """One series as the web app reads it."""
    out: dict[str, Json] = {
        "id": series.id,
        "label": series.label,
        "method": series.method,
        "source": series.source,
        "url": series.url,
        "licence": series.licence,
        "ages": list[Json](series.ages),
        "level": list[Json](series.level),
    }
    if series.note is not None:
        out["note"] = series.note
    if series.low is not None and series.high is not None:
        out["low"] = list[Json](series.low)
        out["high"] = list[Json](series.high)
    return out


def write_sealevel(series: list[Series], folder: Path) -> int:
    """Write every series to ``index.json``; return how many points they hold."""
    folder.mkdir(parents=True, exist_ok=True)
    body = json.dumps(
        {"schemaVersion": 1, "series": [series_json(each) for each in series]},
        separators=(",", ":"),
    )
    _ = (folder / "index.json").write_text(body + "\n")
    return sum(len(each.ages) for each in series)


def _download(source: Source, what: str) -> bytes:  # pragma: no cover - needs network access
    # Each URL is one of the pinned constants above, all https.
    request = urllib.request.Request(source.url, headers={"User-Agent": "histeclipse-pipeline"})  # noqa: S310
    with urllib.request.urlopen(request, timeout=120) as response:  # noqa: S310  # pyright: ignore[reportAny]
        body: object = response.read()  # pyright: ignore[reportAny]
    if not isinstance(body, bytes):
        msg = f"{source.url}: expected bytes"
        raise TypeError(msg)
    return verified(body, source.sha256, what)


def fetch_sealevel(folder: Path) -> int:  # pragma: no cover - needs network access
    """Download the four sources (each checked against its pinned checksum) and write them."""
    return write_sealevel(
        [
            van_der_meer(_download(VAN_DER_MEER, "van der Meer et al. 2022 table")),
            miller(_download(MILLER, "Miller et al. 2020 curve").decode()),
            bintanja(_download(BINTANJA, "Bintanja & van de Wal 2008 file").decode()),
            spratt(_download(SPRATT, "Spratt & Lisiecki 2016 file").decode()),
        ],
        folder,
    )
